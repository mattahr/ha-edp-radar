"""openpyxl helpers shared by the NATO, EDA and SIPRI parsers (S11).

Workbooks are opened read-only with cached values (``data_only``). Cells keep
their font so SIPRI's blue (estimate) and red (highly uncertain) markers can be
read; the legacy indexed palette maps blue to 12 and red to 10.
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from .providers.base import SchemaChangedError

_YEAR = re.compile(r"^(\d{4})\s*(e?)$")
_MISSING = frozenset({"", "..", "...", ". .", "xxx", "-", ":", "n/a", "na"})
_RGB_TO_INDEXED = {"0000FF": 12, "FF0000": 10, "000000": 8}


def open_workbook(payload: bytes) -> Workbook:
    try:
        return load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as err:  # openpyxl raises several unrelated types
        raise SchemaChangedError(f"workbook cannot be opened: {err}") from err


def require_sheet(book: Workbook, name: str) -> Any:
    if name not in book.sheetnames:
        raise SchemaChangedError(f"sheet {name!r} missing; found {book.sheetnames}")
    return book[name]


def rows_of(sheet: Any) -> list[tuple[Any, ...]]:
    return [tuple(row) for row in sheet.iter_rows()]


def text(cell: Any) -> str:
    value = getattr(cell, "value", cell)
    return "" if value is None else str(value).strip()


def number(value: Any) -> Decimal | None:
    """Numeric cell → Decimal; SIPRI/NATO/EDA missing markers → None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return Decimal(str(value))
    if isinstance(value, Decimal):
        return value
    cleaned = str(value).strip().replace(" ", "").replace(" ", "")
    if cleaned.casefold() in _MISSING:
        return None
    try:
        return Decimal(cleaned.replace(",", "."))
    except InvalidOperation:
        return None


def find_row(
    rows: list[tuple[Any, ...]],
    predicate: Callable[[str], bool],
    *,
    start: int = 0,
    what: str,
) -> int:
    for index in range(start, len(rows)):
        if rows[index] and predicate(text(rows[index][0])):
            return index
    raise SchemaChangedError(f"{what} not found from row {start + 1}")


def year_header(row: tuple[Any, ...]) -> dict[int, tuple[int, bool]]:
    """``{column: (year, is_estimate)}`` for cells like ``2024``/``"2025e"``."""
    years: dict[int, tuple[int, bool]] = {}
    for column, cell in enumerate(row):
        value = getattr(cell, "value", cell)
        if isinstance(value, int | float) and not isinstance(value, bool):
            if 1900 <= int(value) <= 2100:
                years[column] = (int(value), False)
            continue
        match = _YEAR.match(text(value))
        if match:
            years[column] = (int(match.group(1)), match.group(2) == "e")
    return years


_BASE_YEAR = re.compile(r"\bconstant\D{0,3}(\d{4})", re.IGNORECASE)


def constant_base_year(label: str) -> int:
    """``2024`` from ``"Constant (2024) US$"`` or ``"constant 2021 prices"``."""
    match = _BASE_YEAR.search(label)
    if match is None:
        raise SchemaChangedError(f"no constant-price base year in {label!r}")
    return int(match.group(1))


def sheet_by_prefix(book: Workbook, prefix: str) -> str:
    matches = [str(name) for name in book.sheetnames if str(name).startswith(prefix)]
    if not matches:
        raise SchemaChangedError(
            f"no sheet starting with {prefix!r}; found {book.sheetnames}"
        )
    if len(matches) > 1:
        raise SchemaChangedError(
            f"{len(matches)} sheets start with {prefix!r}: {matches}"
        )
    return matches[0]


def font_colour_index(cell: Any) -> int | None:
    font = getattr(cell, "font", None)
    colour = getattr(font, "color", None)
    if colour is None:
        return None
    if colour.type == "indexed":
        return int(colour.indexed)
    if colour.type == "rgb" and isinstance(colour.rgb, str):
        return _RGB_TO_INDEXED.get(colour.rgb[-6:].upper())
    return None
