"""openpyxl helper behaviour on a generated workbook (S11)."""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from custom_components.edp_radar.spending.providers.base import SchemaChangedError
from custom_components.edp_radar.spending.xlsx import (
    find_row,
    font_colour_index,
    number,
    open_workbook,
    require_sheet,
    rows_of,
    text,
    year_header,
)


def _workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Table 1"
    sheet.append(["Table 1: Core defence expenditure"])
    sheet.append([None, 2024, "2025e", "2026e"])
    sheet.append(["Sweden (Kronor)", 140549, 187790.5, "..."])
    sheet.cell(row=3, column=3).font = Font(color="FF0000FF")
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_sheet_rows_text_and_numbers() -> None:
    book = open_workbook(_workbook())
    rows = rows_of(require_sheet(book, "Table 1"))
    assert text(rows[0][0]) == "Table 1: Core defence expenditure"
    assert text(rows[1][0]) == ""
    assert number(rows[2][1].value) == Decimal("140549")
    assert number(rows[2][2].value) == Decimal("187790.5")
    assert number("...") is None
    assert number("xxx") is None
    assert number("-") is None
    assert number("1 234,5") == Decimal("1234.5")
    assert number(None) is None
    with pytest.raises(SchemaChangedError, match="Table 9"):
        require_sheet(book, "Table 9")


def test_find_row_and_year_header() -> None:
    rows = rows_of(require_sheet(open_workbook(_workbook()), "Table 1"))
    assert find_row(rows, lambda t: t.startswith("Table 1:"), what="title") == 0
    with pytest.raises(SchemaChangedError, match="subtitle"):
        find_row(rows, lambda t: t == "nope", what="subtitle")
    assert year_header(rows[1]) == {1: (2024, False), 2: (2025, True), 3: (2026, True)}
    assert year_header(rows[0]) == {}


def test_font_colour_index_reads_rgb_and_indexed() -> None:
    rows = rows_of(require_sheet(open_workbook(_workbook()), "Table 1"))
    assert font_colour_index(rows[2][2]) == 12  # blue, mapped from FF0000FF
    assert font_colour_index(rows[2][1]) is None


def test_constant_base_year_and_sheet_prefix() -> None:
    from openpyxl import Workbook

    from custom_components.edp_radar.spending.xlsx import (
        constant_base_year,
        sheet_by_prefix,
    )

    assert constant_base_year("Constant (2024) US$") == 2024
    assert constant_base_year("constant 2021 prices and exchange rates") == 2021
    with pytest.raises(SchemaChangedError, match="base year"):
        constant_base_year("Current US$")
    with pytest.raises(SchemaChangedError):
        constant_base_year("inconstant 1999")
    book = Workbook()
    book.active.title = "Constant (2025) US$"
    assert sheet_by_prefix(book, "Constant (") == "Constant (2025) US$"
    with pytest.raises(SchemaChangedError, match="Share of"):
        sheet_by_prefix(book, "Share of")
    book.create_sheet("Constant (2024) US$")
    with pytest.raises(SchemaChangedError, match="2 sheets"):
        sheet_by_prefix(book, "Constant (")
