"""NATO discovery and workbook parsing (S11; plan §27–§33, §83)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from homeassistant.core import HomeAssistant
from homeassistant.helpers.aiohttp_client import async_get_clientsession
from pytest_homeassistant_custom_component.test_util.aiohttp import (
    AiohttpClientMocker,
)

from custom_components.edp_radar.spending.models import DatapointStatus, SourceRelease
from custom_components.edp_radar.spending.providers.base import SchemaChangedError
from custom_components.edp_radar.spending.providers.nato import (
    TOPIC_URL,
    NatoProvider,
    discover_workbook,
    parse_nato_workbook,
)

FIXTURES = Path(__file__).parent.parent / "fixtures" / "spending" / "nato"
XLSX_2026 = "https://www.nato.int/content/dam/nato/webready/documents/finance/def-exp-2026-en.xlsx"


def _release(year: int = 2026) -> SourceRelease:
    return SourceRelease(
        source_id="nato",
        release_id=f"{year}:test",
        published_at=date(2026, 7, 10),
        download_url=XLSX_2026.replace("2026", str(year)),
        canonical_url=TOPIC_URL,
        format="xlsx",
    )


def test_discover_picks_newest_archive_year() -> None:
    html = (FIXTURES / "topic-page.html").read_text(encoding="utf-8")
    assert discover_workbook(html, TOPIC_URL) == (2026, XLSX_2026)
    with pytest.raises(SchemaChangedError, match="def-exp"):
        discover_workbook("<html></html>", TOPIC_URL)


def test_parse_2026_workbook_sweden_and_status() -> None:
    result = parse_nato_workbook(
        (FIXTURES / "def-exp-2026-en.xlsx").read_bytes(), _release()
    )
    points = {
        (p.metric_id, p.country, p.reference.start.year): p for p in result.datapoints
    }
    assert points[("defence_expenditure_nac", "SE", 2024)].value == Decimal("140549")
    assert points[("defence_expenditure_nac", "SE", 2024)].flags == ("currency:Kronor",)
    assert points[("defence_expenditure_usd_current", "SE", 2026)].value == Decimal(
        "24186"
    )
    assert (
        points[("defence_expenditure_usd_current", "SE", 2026)].status
        is DatapointStatus.ESTIMATE
    )
    assert (
        points[("defence_expenditure_usd_current", "SE", 2024)].status
        is DatapointStatus.ACTUAL
    )
    assert points[("defence_expenditure_usd_constant", "SE", 2026)].value == Decimal(
        "21538"
    )
    assert (
        points[("defence_expenditure_usd_constant", "SE", 2026)].unit
        == "USD_MILLION_CONSTANT_2021"
    )
    assert points[("defence_expenditure_pct_gdp", "SE", 2026)].value == Decimal("3.22")
    assert points[("equipment_share_pct", "SE", 2026)].value == Decimal("25.41")
    derived = points[("equipment_expenditure_usd_current", "SE", 2026)]
    assert derived.value == Decimal("24186") * Decimal("25.41") / Decimal(100)
    assert derived.status is DatapointStatus.ESTIMATE
    assert derived.flags == ("derived",)
    assert points[("defence_expenditure_usd_current", "SI", 2026)].flags == (
        "footnote_star",
    )
    assert points[("defence_expenditure_usd_current", "US", 2026)].value == Decimal(
        "1032849"
    )
    assert points[("defence_expenditure_usd_current", "TR", 2026)].country == "TR"
    countries = {
        p.country
        for p in result.datapoints
        if p.metric_id == "defence_expenditure_usd_current"
    }
    assert len(countries) == 31 and "IS" not in countries
    assert not any(p.country.startswith("NATO") for p in result.datapoints)
    assert result.warnings == ()
    assert result.layout_fingerprint.startswith("Table 1: Core defence expenditure")


def test_parse_2025_workbook_layout_variant() -> None:
    result = parse_nato_workbook(
        (FIXTURES / "def-exp-2025-en.xlsx").read_bytes(), _release(2025)
    )
    points = {
        (p.metric_id, p.country, p.reference.start.year): p for p in result.datapoints
    }
    assert (
        points[("defence_expenditure_usd_current", "SE", 2025)].status
        is DatapointStatus.ESTIMATE
    )
    assert (
        points[("defence_expenditure_usd_current", "SE", 2023)].status
        is DatapointStatus.ACTUAL
    )
    assert points[("defence_expenditure_nac", "SE", 2025)].flags == (
        "currency:Kronor",
        "footnote_star",
    )
    assert points[("equipment_share_pct", "SE", 2025)].value == Decimal(
        "35.82725060827251"
    )
    assert result.layout_fingerprint.startswith("Table 1: Defence expenditure")


def test_base_year_comes_from_the_table_subtitle() -> None:
    import io

    import openpyxl

    book = openpyxl.load_workbook(FIXTURES / "def-exp-2026-en.xlsx")
    sheet = book["Table 2"]
    target = None
    for row in sheet.iter_rows():
        for cell in row[:4]:
            value = cell.value
            if isinstance(value, str) and "onstant 2021" in value.casefold():
                target = cell
                break
        if target is not None:
            break
    assert target is not None, "no cell mentioning 'onstant 2021' found in Table 2"
    target.value = target.value.replace("2021", "2019")
    buffer = io.BytesIO()
    book.save(buffer)
    result = parse_nato_workbook(buffer.getvalue(), _release())
    constant_points = [
        p
        for p in result.datapoints
        if p.metric_id == "defence_expenditure_usd_constant"
    ]
    assert constant_points
    assert all(p.unit == "USD_MILLION_CONSTANT_2019" for p in constant_points)


def test_layout_change_fails() -> None:
    import io

    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Table 1"
    book.active.append(["Table 1: Something else"])
    for name in ("Table 2", "Table 3", "Table 8a"):
        book.create_sheet(name).append([f"{name}: x"])
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(SchemaChangedError):
        parse_nato_workbook(buffer.getvalue(), _release())


async def test_provider_discovery_uses_head_metadata(
    hass: HomeAssistant, aioclient_mock: AiohttpClientMocker
) -> None:
    aioclient_mock.get(
        TOPIC_URL, text=(FIXTURES / "topic-page.html").read_text(encoding="utf-8")
    )
    aioclient_mock.head(
        XLSX_2026,
        headers={
            "ETag": '"0x8DEDE695735E099"',
            "Last-Modified": "Fri, 10 Jul 2026 09:55:14 GMT",
        },
    )
    provider = NatoProvider()
    session = async_get_clientsession(hass)
    release = await provider.async_discover_latest(session)
    assert release.release_id == '2026:"0x8DEDE695735E099"'
    assert release.published_at == date(2026, 7, 10)
    assert release.download_url == XLSX_2026
    aioclient_mock.get(
        XLSX_2026, content=(FIXTURES / "def-exp-2026-en.xlsx").read_bytes()
    )
    fetched, payload = await provider.async_fetch_release(session, release)
    assert fetched.checksum
    assert len(provider.parse_release(payload, fetched).datapoints) > 1500
