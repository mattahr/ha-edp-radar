"""SIPRI discovery and workbook parsing (S11; plan §40–§44, §85)."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from homeassistant.core import HomeAssistant
from homeassistant.helpers.aiohttp_client import async_get_clientsession
from pytest_homeassistant_custom_component.test_util.aiohttp import (
    AiohttpClientMocker,
)

from custom_components.edp_radar.spending.models import DatapointStatus, SourceRelease
from custom_components.edp_radar.spending.providers.base import SchemaChangedError
from custom_components.edp_radar.spending.providers.sipri import (
    LANDING_URL,
    MIN_YEAR,
    SipriProvider,
    discover_release,
    parse_sipri_workbook,
)

FIXTURES = Path(__file__).parent.parent / "fixtures" / "spending" / "sipri"
XLSX_URL = (
    "https://www.sipri.org/sites/default/files/SIPRI-Milex-data-1949-2025_v1.2.xlsx"
)


def _release() -> SourceRelease:
    return SourceRelease(
        "sipri",
        "SIPRI-Milex-data-1949-2025_v1.2.xlsx:2026-04-27",
        date(2026, 4, 27),
        XLSX_URL,
        LANDING_URL,
        "xlsx",
    )


def test_discover_link_and_revision_date() -> None:
    html = (FIXTURES / "landing.html").read_text(encoding="utf-8")
    assert discover_release(html, LANDING_URL) == (XLSX_URL, date(2026, 4, 27))
    assert discover_release(
        '<a href="https://www.sipri.org/x/y.xlsx">x</a>', LANDING_URL
    ) == ("https://www.sipri.org/x/y.xlsx", None)
    with pytest.raises(SchemaChangedError, match="xlsx"):
        discover_release("<html></html>", LANDING_URL)


def test_parse_sweden_estimates_and_notes() -> None:
    result = parse_sipri_workbook(
        (FIXTURES / "milex-trimmed.xlsx").read_bytes(), _release()
    )
    points = {
        (p.metric_id, p.country, p.reference.start.year): p for p in result.datapoints
    }
    assert points[("military_expenditure_usd_constant", "SE", 2025)].value == Decimal(
        "14954.07135864315"
    )
    assert points[("military_expenditure_usd_current", "SE", 2025)].value == Decimal(
        "16473.47720649671"
    )
    assert (
        points[("military_expenditure_pct_gdp", "SE", 2025)].value
        == Decimal("0.02471184849540041") * 100
    )
    assert (
        points[("military_expenditure_usd_constant", "SE", 2025)].status
        is DatapointStatus.ACTUAL
    )
    assert (
        points[("military_expenditure_usd_constant", "SE", 1990)].status
        is DatapointStatus.ACTUAL
    )
    assert (
        "military_expenditure_usd_constant",
        "SE",
        1975,
    ) not in points  # below MIN_YEAR
    assert MIN_YEAR == 1990
    albania = points[("military_expenditure_usd_constant", "AL", 2025)]
    assert albania.status is DatapointStatus.BUDGET
    assert albania.flags == ("excludes_paramilitary", "footnote_3")
    assert points[("military_expenditure_usd_constant", "IS", 2025)].value == Decimal(
        "0"
    )
    assert points[("military_expenditure_usd_current", "TR", 2025)].flags == (
        "currency_redenominated",
        "footnote_105",
    )
    assert not any(p.country in {"USSR", "European Union"} for p in result.datapoints)
    assert result.warnings == ()
    assert (
        len(
            {
                p.country
                for p in result.datapoints
                if p.metric_id == "military_expenditure_usd_current"
            }
        )
        == 59
    )
    assert "blue" in result.layout_fingerprint.casefold()


def test_blue_font_is_estimate_when_present_above_min_year() -> None:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    for name in ("Constant (2024) US$", "Current US$", "Share of GDP"):
        sheet = book.create_sheet(name)
        sheet.append(["Military expenditure by country"])
        sheet.append(["x"])
        sheet.append(
            [
                "Figures in blue are SIPRI estimates. Figures in red indicate "
                "highly uncertain data."
            ]
        )
        sheet.append(["Country", "Notes", 1990, 1991])
        sheet.append(["Europe"])
        sheet.append(["Sweden", None, 100, 200])
        sheet.cell(row=6, column=3).font = Font(color="FF0000FF")
        sheet.cell(row=6, column=4).font = Font(color="FFFF0000")
    book.remove(book["Sheet"])
    buffer = io.BytesIO()
    book.save(buffer)
    result = parse_sipri_workbook(buffer.getvalue(), _release())
    points = {(p.metric_id, p.reference.start.year): p for p in result.datapoints}
    assert (
        points[("military_expenditure_usd_constant", 1990)].status
        is DatapointStatus.ESTIMATE
    )
    assert (
        points[("military_expenditure_usd_constant", 1991)].status
        is DatapointStatus.ACTUAL
    )
    assert points[("military_expenditure_usd_constant", 1991)].flags == (
        "highly_uncertain",
    )


def test_missing_sheet_or_legend_fails() -> None:
    import io

    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Constant (2024) US$"
    book.active.append(["Country", "Notes", 2020])
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(SchemaChangedError):
        parse_sipri_workbook(buffer.getvalue(), _release())


async def test_provider_round_trip(
    hass: HomeAssistant, aioclient_mock: AiohttpClientMocker
) -> None:
    aioclient_mock.get(
        LANDING_URL, text=(FIXTURES / "landing.html").read_text(encoding="utf-8")
    )
    aioclient_mock.head(
        XLSX_URL,
        headers={
            "ETag": '"e13b8-65072a76c0127"',
            "Last-Modified": "Mon, 27 Apr 2026 15:20:25 GMT",
        },
    )
    provider = SipriProvider()
    session = async_get_clientsession(hass)
    release = await provider.async_discover_latest(session)
    assert release.release_id == "SIPRI-Milex-data-1949-2025_v1.2.xlsx:2026-04-27"
    assert release.published_at == date(2026, 4, 27)
    aioclient_mock.get(XLSX_URL, content=(FIXTURES / "milex-trimmed.xlsx").read_bytes())
    fetched, payload = await provider.async_fetch_release(session, release)
    assert fetched.etag == '"e13b8-65072a76c0127"'
    assert len(provider.parse_release(payload, fetched).datapoints) > 5000
