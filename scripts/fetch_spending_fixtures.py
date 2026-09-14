"""Download and trim the real spending fixtures (S18).

Usage:
    PYTHONPATH=. uv run python scripts/fetch_spending_fixtures.py statskontoret
    PYTHONPATH=. uv run python scripts/fetch_spending_fixtures.py all

Downloads are cached under --cache (default .cache/spending/fixtures). Trimming
keeps the rows/sheets the parsers and tests need; numbers are never edited.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import io
import re
import sys
import zipfile
import zlib
from pathlib import Path
from urllib.parse import quote
from urllib.request import Request, urlopen

import openpyxl

FIXTURES = Path("tests/fixtures/spending")
USER_AGENT = "ha-edp-radar fixture fetcher"


def download(url: str, cache: Path, name: str) -> bytes:
    cache.mkdir(parents=True, exist_ok=True)
    target = cache / name
    if target.exists():
        return target.read_bytes()
    # Python 3.14's http.client defaults to "Accept-Encoding: identity"; the
    # Statskontoret server 404s on that header, so request gzip/deflate explicitly
    # and decompress the response ourselves (http.client does not do it for us).
    headers = {"User-Agent": USER_AGENT, "Accept-Encoding": "gzip, deflate"}
    with urlopen(Request(url, headers=headers)) as response:
        data: bytes = response.read()
        encoding = response.headers.get("Content-Encoding", "")
    if encoding == "gzip":
        data = gzip.decompress(data)
    elif encoding == "deflate":
        data = zlib.decompress(data)
    target.write_bytes(data)
    print(f"downloaded {url} -> {target} ({len(data)} bytes)")
    return data


def write(path: Path, data: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    print(f"wrote {path} ({len(data)} bytes)")


# ----------------------------------------------------------------- Statskontoret

SK_BASE = "https://www.statskontoret.se"
SK_PAGE = SK_BASE + "/analys-och-statistik/oppna-data/manadsutfall/?year={year}"
SK_FILE = (
    SK_BASE
    + "/OpenDataManadsUtfallPage/GetFile?documentType=Utgift&fileType=Zip"
    + "&fileName=x.zip&Year={year}&month={month}&status={status}"
)
_LI_DATA = re.compile(r'<li class="data[^"]*">.*?</li>', re.S)


def trim_statskontoret_page(html: str) -> str:
    """Keep only the ``<li class="data">`` blocks for Utgifter releases."""
    blocks = [b for b in _LI_DATA.findall(html) if "documentType=Utgift" in b]
    body = "\n".join(blocks)
    return (
        f'<!DOCTYPE html>\n<html lang="sv"><body><ul class="fixture">\n'
        f"{body}\n</ul></body></html>\n"
    )


def trim_statskontoret_csv(zip_bytes: bytes, years: set[str]) -> bytes:
    """Header + every utgiftsområde 06 row of ``years`` + three non-06 rows."""
    archive = zipfile.ZipFile(io.BytesIO(zip_bytes))
    text = archive.read(archive.namelist()[0]).decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    kept = [rows[0]]
    others = 0
    for row in rows[1:]:
        if row[0] == "06" and row[10] in years:
            kept.append(row)
        elif others < 3 and row[10] in years:
            kept.append(row)
            others += 1
    out = io.StringIO()
    csv.writer(out, delimiter=";", lineterminator="\n").writerows(kept)
    return ("﻿" + out.getvalue()).encode("utf-8")


def fetch_statskontoret(cache: Path) -> None:
    folder = FIXTURES / "statskontoret"
    for year in (2026, 2025):
        page = download(SK_PAGE.format(year=year), cache, f"statskontoret-{year}.html")
        write(
            folder / f"discovery-{year}.html",
            trim_statskontoret_page(page.decode("utf-8")).encode("utf-8"),
        )
    july = download(
        SK_FILE.format(year=2026, month=7, status="Definitiv"),
        cache,
        "statskontoret-2026-07.zip",
    )
    write(
        folder / "utgifter-2026-07.csv", trim_statskontoret_csv(july, {"2025", "2026"})
    )
    prelim = download(
        SK_FILE.format(year=2025, month=12, status=quote("Preliminär 1")),
        cache,
        "statskontoret-2025-12-preliminar.zip",
    )
    write(
        folder / "utgifter-2025-12-preliminar.csv",
        trim_statskontoret_csv(prelim, {"2025"}),
    )
    definitive = download(
        SK_FILE.format(year=2025, month=12, status="Definitiv"),
        cache,
        "statskontoret-2025-12-definitiv.zip",
    )
    write(
        folder / "utgifter-2025-12-definitiv.csv",
        trim_statskontoret_csv(definitive, {"2025"}),
    )


# ---------------------------------------------------------------------- Eurostat

EUROSTAT_URL = (
    "https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/gov_ev"
    "?lang=en&expend=DEF&na_item=TE&na_item=P51G&unit=MIO_EUR&unit=PC_GDP&unit=MIO_NAC"
)


def fetch_eurostat(cache: Path) -> None:
    payload = download(EUROSTAT_URL, cache, "eurostat-gov_ev.json")
    write(FIXTURES / "eurostat" / "gov_ev-defence.json", payload)


# -------------------------------------------------------------------------- NATO

NATO_TOPIC = "https://www.nato.int/en/what-we-do/introduction-to-nato/defence-expenditures-and-natos-5-commitment"
NATO_XLSX = "https://www.nato.int/content/dam/nato/webready/documents/finance/def-exp-{year}-en.xlsx"
_NATO_ARCHIVE = re.compile(
    r'<a href=\\?"[^"]*def-exp-\d{4}-en\.(?:pdf|PDF)\\?"[^>]*>\d{4}</a>'
)


def trim_nato_topic(html: str) -> str:
    """Keep only the archive anchors (``def-exp-YYYY-en.pdf``) the discovery needs."""
    anchors = _NATO_ARCHIVE.findall(html)
    body = "\n".join(a.replace('\\"', '"') for a in anchors)
    return (
        f'<!DOCTYPE html>\n<html lang="en"><body><p>Archive of tables</p>\n'
        f"{body}\n</body></html>\n"
    )


def fetch_nato(cache: Path) -> None:
    folder = FIXTURES / "nato"
    topic = download(NATO_TOPIC, cache, "nato-topic.html")
    write(
        folder / "topic-page.html",
        trim_nato_topic(topic.decode("utf-8", "replace")).encode("utf-8"),
    )
    for year in (2026, 2025):
        write(
            folder / f"def-exp-{year}-en.xlsx",
            download(NATO_XLSX.format(year=year), cache, f"def-exp-{year}-en.xlsx"),
        )


# --------------------------------------------------------------------------- EDA

EDA_PORTAL = "https://www.eda.europa.eu/publications-and-data/defence-data"
_EDA_ANCHOR = re.compile(
    r"""<a[^>]*href=(?:"|')[^"']*\.xlsx(?:"|')[^>]*>\s*(?:<span>)?\s*"""
    r"""Defence Data \d{4}\s*(?:</span>)?\s*</a>""",
    re.I | re.S,
)
_EDA_LINK = re.compile(
    r"""href=(["'])([^"']*\.xlsx)\1[^>]*>\s*(?:<span>)?\s*Defence Data (\d{4})""",
    re.I | re.S,
)


def trim_eda_portal(html: str) -> str:
    anchors = _EDA_ANCHOR.findall(html)
    return (
        '<!DOCTYPE html>\n<html lang="en"><body>\n'
        + "\n".join(anchors)
        + "\n</body></html>\n"
    )


def fetch_eda(cache: Path) -> None:
    folder = FIXTURES / "eda"
    portal = download(EDA_PORTAL, cache, "eda-portal.html").decode("utf-8", "replace")
    write(folder / "portal.html", trim_eda_portal(portal).encode("utf-8"))
    links = {int(year): url for _, url, year in _EDA_LINK.findall(portal)}
    for year in (2025, 2022):
        write(
            folder / f"defence-data-{year}.xlsx",
            download(links[year], cache, f"eda-{year}.xlsx"),
        )


# ------------------------------------------------------------------------- SIPRI

SIPRI_LANDING = "https://www.sipri.org/databases/milex"
_SIPRI_XLSX = re.compile(r'href="((?:https?:)?//www\.sipri\.org/[^"]*\.xlsx)"')
_SIPRI_REVISED = re.compile(r"revised on \d{1,2} [A-Za-z]+ \d{4}[^.<]{0,200}")
SIPRI_SHEETS = ("Constant (2024) US$", "Current US$", "Share of GDP")


def trim_sipri_landing(html: str) -> str:
    link = _SIPRI_XLSX.search(html)
    revised = _SIPRI_REVISED.search(html)
    if link is None or revised is None:
        raise SystemExit("SIPRI landing page: xlsx link or revision sentence not found")
    return (
        '<!DOCTYPE html>\n<html lang="en"><body>\n'
        f"<p>The file was {revised.group(0)}.</p>\n"
        f'<a href="{link.group(1)}">'
        "Download the SIPRI Military Expenditure Database (Excel)</a>\n"
        "</body></html>\n"
    )


KEEP_FROM = "Europe"
KEEP_ROW = "Libya"  # a red-font (highly uncertain) row outside Europe


def trim_sipri_workbook(data: bytes) -> bytes:
    """Keep three data sheets, the rows from 'Europe' on, and one African row
    with red-font cells so the highly-uncertain marker is tested against real
    data (colours survive)."""
    book = openpyxl.load_workbook(io.BytesIO(data))
    for name in list(book.sheetnames):
        if name not in SIPRI_SHEETS:
            book.remove(book[name])
    for name in SIPRI_SHEETS:
        sheet = book[name]
        header = next(
            i for i in range(1, 15) if str(sheet.cell(i, 1).value).strip() == "Country"
        )
        rows = {
            str(sheet.cell(i, 1).value).strip(): i
            for i in range(header + 1, sheet.max_row + 1)
        }
        europe, keep = rows[KEEP_FROM], rows[KEEP_ROW]
        # Delete from the bottom up so earlier indexes stay valid.
        sheet.delete_rows(keep + 1, europe - (keep + 1))
        sheet.delete_rows(header + 2, keep - (header + 2))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def fetch_sipri(cache: Path) -> None:
    folder = FIXTURES / "sipri"
    landing = download(SIPRI_LANDING, cache, "sipri-landing.html").decode(
        "utf-8", "replace"
    )
    write(folder / "landing.html", trim_sipri_landing(landing).encode("utf-8"))
    url = _SIPRI_XLSX.search(landing).group(1)  # type: ignore[union-attr]
    workbook = download(
        "https:" + url if url.startswith("//") else url, cache, "sipri-milex.xlsx"
    )
    write(folder / "milex-trimmed.xlsx", trim_sipri_workbook(workbook))


FETCHERS = {
    "statskontoret": fetch_statskontoret,
    "eurostat": fetch_eurostat,
    "nato": fetch_nato,
    "eda": fetch_eda,
    "sipri": fetch_sipri,
}


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("source", choices=[*FETCHERS, "all"])
    parser.add_argument("--cache", default=".cache/spending/fixtures")
    args = parser.parse_args(argv)
    cache = Path(args.cache)
    for name, fetcher in FETCHERS.items():
        if args.source in (name, "all"):
            fetcher(cache)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
